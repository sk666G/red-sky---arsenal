# language: Python, file: Program/report/deliverable.py, target: Red Sky report — deliverable renderer
import argparse
import html
import json
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Dict, List

from Program.theme.palette import SCARLET, ARTERY, BONE, ASH, OK, CLOT, RESET, BOLD
from Program.utils import print_ok, print_err, print_info, print_warn, print_kv
from Program.utils.paths import OUTPUT_DIR


REPORT_DIR = OUTPUT_DIR / "report"
DELIV_DIR = REPORT_DIR / "deliverables"
DELIV_DIR.mkdir(parents=True, exist_ok=True)
FINDINGS_DIR = OUTPUT_DIR / "findings"
EXEC_SUMMARY_FILE = REPORT_DIR / "exec_summary.md"


SEVERITY_ORDER = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
SEVERITY_ICON = {"critical": "[CRIT]", "high": "[HIGH]", "medium": "[MED]",
                 "low": "[LOW]", "info": "[INFO]"}
SEVERITY_CLASS = {"critical": "crit", "high": "high", "medium": "med",
                  "low": "low", "info": "info"}


def _load_findings() -> List[Dict]:
    out = []
    if not FINDINGS_DIR.exists():
        return out
    for p in sorted(FINDINGS_DIR.glob("*.json")):
        try:
            d = json.loads(p.read_text())
            if isinstance(d, list):
                out.extend(d)
            else:
                out.append(d)
        except (OSError, json.JSONDecodeError):
            continue
    return sorted(out, key=lambda f: (SEVERITY_ORDER.get(f.get("severity", "medium"), 9),
                                      -float(f.get("cvss") or 0)))


def _load_exec_summary() -> str:
    if EXEC_SUMMARY_FILE.exists():
        return EXEC_SUMMARY_FILE.read_text()
    return ""


MD_TEMPLATE = """# {title}

**Client:** {client}
**Engagement:** {engagement}
**Date range:** {date_start} -> {date_end}
**Prepared by:** {author}
**Report ID:** {report_id}

---

## Executive Summary

{exec_summary}

---

## Scope & Methodology

**In-scope targets:**

{scope}

**Methodology:** Testing followed industry-standard phases: reconnaissance,
enumeration, exploitation, post-exploitation, and reporting. Findings are
severity-rated using CVSS 3.1 base scores and business-impact weighting.

---

## Findings Summary

| # | ID | Severity | CVSS | Title | Target |
|---|-----|----------|------|-------|--------|
{findings_table}

**Totals:** {total} findings - {critical_count} critical, {high_count} high, {medium_count} medium, {low_count} low, {info_count} info.

---

## Findings Detail

{findings_detail}

---

## Remediation Roadmap

{roadmap}

---

## Appendix

### A. Evidence Files

{evidence_files}

### B. Tooling

{tooling}

### C. Raw Findings

{raw_json}
"""


FINDING_MD = """### {icon} [{id}] {title}

**Severity:** {severity_cap} | **CVSS:** {cvss} | **Target:** `{target}`
**Module:** {module} | **CWE:** {cwe}

**Evidence:**

{evidence}

**Remediation:**

{remediation}

**References:**

{references}

---
"""


ROADMAP = """**Immediate (0-7 days):**
{immediate}

**Short term (1-30 days):**
{short}

**Medium term (30-90 days):**
{medium}

**Long term (90+ days):**
{long}
"""


def _render_md(meta: Dict) -> str:
    findings = _load_findings()
    counts = {s: 0 for s in SEVERITY_ORDER}
    for f in findings:
        sev = f.get("severity", "medium")
        counts[sev] = counts.get(sev, 0) + 1

    rows = []
    for i, f in enumerate(findings, 1):
        rows.append("| {} | {} | {} {} | {} | {} | `{}` |".format(
            i, f.get("id", ""), SEVERITY_ICON.get(f.get("severity", ""), ""),
            f.get("severity", "").upper(), f.get("cvss", ""),
            (f.get("title") or "").replace("|", "\\\\|"),
            (f.get("target") or "").replace("|", "\\\\|")))

    detail = []
    for f in findings:
        evidence = "\\n".join("- `" + str(e) + "`" for e in (f.get("evidence") or ["(none provided)"]))
        refs = "\\n".join("- " + str(r) for r in (f.get("references") or ["(none)"]))
        detail.append(FINDING_MD.format(
            icon=SEVERITY_ICON.get(f.get("severity", ""), "*"),
            id=f.get("id", ""),
            title=f.get("title", "untitled"),
            severity_cap=f.get("severity", "medium").upper(),
            cvss=f.get("cvss", 0),
            target=f.get("target", ""),
            module=f.get("module", ""),
            cwe=f.get("cwe", "(not set)"),
            evidence=evidence,
            remediation=f.get("remediation", "(not provided)"),
            references=refs,
        ))

    imm, short, med, long_ = [], [], [], []
    for f in findings:
        sev = f.get("severity", "medium")
        line = "- **" + (f.get("id") or "") + "** " + (f.get("title") or "")
        if sev == "critical":
            imm.append(line)
        elif sev == "high":
            short.append(line)
        elif sev == "medium":
            med.append(line)
        else:
            long_.append(line)

    tooling = [
        "**Red Sky Arsenal** - orchestration + module framework",
        "hashcat / john - hash cracking",
        "nmap / masscan - port discovery",
        "mitmproxy / Burp - web app testing",
        "impacket - Windows protocol abuse",
        "proxmark3 - RFID",
    ]

    raw = json.dumps(findings, indent=2)

    return MD_TEMPLATE.format(
        title=meta["title"],
        client=meta["client"],
        engagement=meta["engagement"],
        date_start=meta["date_start"],
        date_end=meta["date_end"],
        author=meta["author"],
        report_id=meta["report_id"],
        exec_summary=_load_exec_summary() or "(executive summary not yet generated)",
        scope="\\n".join("- `" + t + "`" for t in meta["scope"]) or "(none specified)",
        findings_table="\\n".join(rows) or "| | | | | (no findings) | |",
        total=len(findings),
        critical_count=counts.get("critical", 0),
        high_count=counts.get("high", 0),
        medium_count=counts.get("medium", 0),
        low_count=counts.get("low", 0),
        info_count=counts.get("info", 0),
        findings_detail="\\n".join(detail) or "(no findings)",
        roadmap=ROADMAP.format(
            immediate="\\n".join(imm) or "(none)",
            short="\\n".join(short) or "(none)",
            medium="\\n".join(med) or "(none)",
            long="\\n".join(long_) or "(none)",
        ),
        evidence_files=meta.get("evidence_files") or "(none listed)",
        tooling="\\n".join("- " + t for t in tooling),
        raw_json=raw,
    )


HTML_TEMPLATE = """<!doctype html>
<html><head><meta charset="utf-8"><title>{title}</title>
<style>
  body {{ font-family: -apple-system, system-ui, Segoe UI, sans-serif;
          max-width: 900px; margin: 2em auto; padding: 0 2em; color: #222; line-height: 1.5; }}
  h1 {{ border-bottom: 3px solid #8b0000; padding-bottom: .3em; }}
  h2 {{ margin-top: 2em; color: #8b0000; }}
  h3 {{ margin-top: 1.6em; }}
  table {{ border-collapse: collapse; width: 100%; margin: 1em 0; }}
  th, td {{ border: 1px solid #ddd; padding: .5em .7em; text-align: left; font-size: 14px; }}
  th {{ background: #f4f4f7; }}
  .meta {{ background: #f4f4f7; padding: 1em 1.2em; border-left: 4px solid #8b0000; }}
  .crit {{ color: #8b0000; font-weight: bold; }}
  .high {{ color: #dc143c; font-weight: bold; }}
  .med {{ color: #c2650f; }}
  .low {{ color: #2b6cb0; }}
  .info {{ color: #888; }}
  pre, code {{ background: #f4f4f7; padding: .15em .35em; border-radius: 3px;
               font-family: ui-monospace, Menlo, Consolas, monospace; font-size: 13px; }}
  pre {{ padding: 1em; overflow-x: auto; }}
  .finding {{ border-left: 4px solid #ddd; padding: .3em 0 .3em 1em; margin: 1.4em 0; }}
  .finding.crit {{ border-color: #8b0000; }}
  .finding.high {{ border-color: #dc143c; }}
  .finding.med {{ border-color: #c2650f; }}
  .finding.low {{ border-color: #2b6cb0; }}
  .finding.info {{ border-color: #ccc; }}
</style>
</head><body>
<h1>{title}</h1>
<div class="meta">
  <div><b>Client:</b> {client}</div>
  <div><b>Engagement:</b> {engagement}</div>
  <div><b>Date range:</b> {date_start} &rarr; {date_end}</div>
  <div><b>Prepared by:</b> {author}</div>
  <div><b>Report ID:</b> {report_id}</div>
</div>

<h2>Executive Summary</h2>
<div>{exec_summary_html}</div>

<h2>Scope &amp; Methodology</h2>
<p><b>In-scope:</b></p>
<ul>{scope_html}</ul>
<p>Testing followed industry-standard phases: reconnaissance, enumeration,
exploitation, post-exploitation, and reporting. Findings rated on CVSS 3.1
with business-impact weighting.</p>

<h2>Findings Summary</h2>
<table>
<thead><tr><th>#</th><th>ID</th><th>Severity</th><th>CVSS</th><th>Title</th><th>Target</th></tr></thead>
<tbody>
{findings_table_html}
</tbody></table>
<p><b>Totals:</b> {total} findings - {critical_count} critical, {high_count} high, {medium_count} medium, {low_count} low, {info_count} info.</p>

<h2>Findings Detail</h2>
{findings_detail_html}

<h2>Remediation Roadmap</h2>
{roadmap_html}

<h2>Appendix</h2>
<h3>A. Evidence Files</h3>
{evidence_files_html}
<h3>B. Tooling</h3>
<ul>{tooling_html}</ul>
<h3>C. Raw Findings</h3>
<pre>{raw_json_html}</pre>
</body></html>
"""


def _render_html(meta: Dict) -> str:
    findings = _load_findings()
    counts = {s: 0 for s in SEVERITY_ORDER}
    for f in findings:
        sev = f.get("severity", "medium")
        counts[sev] = counts.get(sev, 0) + 1

    rows = []
    for i, f in enumerate(findings, 1):
        sev = f.get("severity", "medium")
        cls = SEVERITY_CLASS.get(sev, "")
        rows.append(
            "<tr><td>{}</td><td>{}</td><td class='{}'>{}</td><td>{}</td><td>{}</td><td><code>{}</code></td></tr>".format(
                i, html.escape(str(f.get("id", ""))), cls, html.escape(sev.upper()),
                html.escape(str(f.get("cvss", ""))),
                html.escape(str(f.get("title", ""))),
                html.escape(str(f.get("target", "")))))

    detail = []
    for f in findings:
        sev = f.get("severity", "medium")
        cls = SEVERITY_CLASS.get(sev, "")
        evidence = "".join("<li><code>" + html.escape(str(e)) + "</code></li>" for e in (f.get("evidence") or ["(none)"]))
        refs = "".join("<li>" + html.escape(str(r)) + "</li>" for r in (f.get("references") or ["(none)"]))
        detail.append(
            '<div class="finding ' + cls + '">'
            '<h3>' + SEVERITY_ICON.get(sev, "*") + ' [' + html.escape(str(f.get("id", ""))) + '] ' + html.escape(str(f.get("title", ""))) + '</h3>'
            '<p><b>Severity:</b> <span class="' + cls + '">' + html.escape(sev.upper()) + '</span> | <b>CVSS:</b> ' + html.escape(str(f.get("cvss", 0))) + ' | <b>Target:</b> <code>' + html.escape(str(f.get("target", ""))) + '</code></p>'
            '<p><b>Module:</b> ' + html.escape(str(f.get("module", ""))) + ' | <b>CWE:</b> ' + html.escape(str(f.get("cwe", ""))) + '</p>'
            '<p><b>Evidence:</b></p><ul>' + evidence + '</ul>'
            '<p><b>Remediation:</b> ' + html.escape(str(f.get("remediation", ""))) + '</p>'
            '<p><b>References:</b></p><ul>' + refs + '</ul>'
            '</div>'
        )

    ex = _load_exec_summary()
    ex_html = ""
    for line in ex.splitlines():
        if line.startswith("### "):
            ex_html += "<h3>" + html.escape(line[4:]) + "</h3>"
        elif line.startswith("## "):
            ex_html += "<h2>" + html.escape(line[3:]) + "</h2>"
        elif line.startswith("# "):
            ex_html += "<h1>" + html.escape(line[2:]) + "</h1>"
        elif line.startswith("- "):
            ex_html += "<li>" + html.escape(line[2:]) + "</li>"
        elif line.strip():
            ex_html += "<p>" + html.escape(line) + "</p>"
    ex_html = ex_html or "<p>(executive summary not yet generated)</p>"

    imm, short, med, long_ = [], [], [], []
    for f in findings:
        sev = f.get("severity", "medium")
        line = "<li><b>" + html.escape(str(f.get("id", ""))) + "</b> " + html.escape(str(f.get("title", ""))) + "</li>"
        if sev == "critical": imm.append(line)
        elif sev == "high":   short.append(line)
        elif sev == "medium": med.append(line)
        else:                 long_.append(line)
    roadmap_html = (
        "<h3>Immediate (0-7 days)</h3><ul>" + ("".join(imm) or "<li>(none)</li>") + "</ul>"
        "<h3>Short term (1-30 days)</h3><ul>" + ("".join(short) or "<li>(none)</li>") + "</ul>"
        "<h3>Medium term (30-90 days)</h3><ul>" + ("".join(med) or "<li>(none)</li>") + "</ul>"
        "<h3>Long term (90+ days)</h3><ul>" + ("".join(long_) or "<li>(none)</li>") + "</ul>"
    )

    return HTML_TEMPLATE.format(
        title=html.escape(meta["title"]),
        client=html.escape(meta["client"]),
        engagement=html.escape(meta["engagement"]),
        date_start=html.escape(meta["date_start"]),
        date_end=html.escape(meta["date_end"]),
        author=html.escape(meta["author"]),
        report_id=html.escape(meta["report_id"]),
        exec_summary_html=ex_html,
        scope_html="".join("<li><code>" + html.escape(str(t)) + "</code></li>" for t in meta["scope"]) or "<li>(none)</li>",
        findings_table_html="".join(rows) or "<tr><td colspan='6'>(no findings)</td></tr>",
        total=len(findings),
        critical_count=counts.get("critical", 0),
        high_count=counts.get("high", 0),
        medium_count=counts.get("medium", 0),
        low_count=counts.get("low", 0),
        info_count=counts.get("info", 0),
        findings_detail_html="".join(detail) or "<p>(no findings)</p>",
        roadmap_html=roadmap_html,
        evidence_files_html="<pre>" + html.escape(str(meta.get("evidence_files") or "(none)")) + "</pre>",
        tooling_html="".join("<li>" + html.escape(t) + "</li>" for t in [
            "Red Sky Arsenal", "hashcat / john", "nmap / masscan",
            "mitmproxy / Burp", "impacket", "proxmark3"]),
        raw_json_html=html.escape(json.dumps(findings, indent=2)),
    )


def _render_docx(meta: Dict, out_path: Path) -> bool:
    try:
        from docx import Document
        from docx.shared import RGBColor
    except ImportError:
        print_warn("python-docx not installed — saving Markdown instead")
        alt = out_path.with_suffix(".md")
        alt.write_text(_render_md(meta))
        print_kv("saved", alt)
        return False

    findings = _load_findings()
    doc = Document()
    doc.add_heading(meta["title"], 0)

    for label, val in (("Client", meta["client"]), ("Engagement", meta["engagement"]),
                       ("Date range", meta["date_start"] + " -> " + meta["date_end"]),
                       ("Prepared by", meta["author"]), ("Report ID", meta["report_id"])):
        p = doc.add_paragraph()
        p.add_run(label + ": ").bold = True
        p.add_run(val)

    doc.add_heading("Executive Summary", level=1)
    for line in _load_exec_summary().splitlines():
        if line.strip():
            doc.add_paragraph(line)

    doc.add_heading("Scope & Methodology", level=1)
    for t in meta["scope"]:
        doc.add_paragraph(t, style="List Bullet")
    doc.add_paragraph(
        "Testing followed industry-standard phases: reconnaissance, enumeration, "
        "exploitation, post-exploitation, and reporting. Findings rated on CVSS 3.1."
    )

    doc.add_heading("Findings Summary", level=1)
    table = doc.add_table(rows=1, cols=6)
    hdr = table.rows[0].cells
    for i, h in enumerate(["#", "ID", "Severity", "CVSS", "Title", "Target"]):
        hdr[i].text = h
    for i, f in enumerate(findings, 1):
        row = table.add_row().cells
        row[0].text = str(i)
        row[1].text = str(f.get("id", ""))
        row[2].text = str(f.get("severity", "")).upper()
        row[3].text = str(f.get("cvss", ""))
        row[4].text = str(f.get("title", ""))
        row[5].text = str(f.get("target", ""))

    doc.add_heading("Findings Detail", level=1)
    for f in findings:
        sev = f.get("severity", "medium")
        doc.add_heading("[" + str(f.get("id", "")) + "] " + str(f.get("title", "")), level=2)
        p = doc.add_paragraph("Severity: {}  |  CVSS: {}  |  Target: {}".format(
            sev.upper(), f.get("cvss", ""), f.get("target", "")))
        if sev == "critical":
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
                run.bold = True
        elif sev == "high":
            for run in p.runs:
                run.font.color.rgb = RGBColor(0xDC, 0x14, 0x3C)
                run.bold = True

        doc.add_paragraph("Evidence:", style="Heading 3")
        for e in (f.get("evidence") or ["(none)"]):
            doc.add_paragraph(str(e), style="List Bullet")
        doc.add_paragraph("Remediation:", style="Heading 3")
        doc.add_paragraph(str(f.get("remediation", "(not provided)")))
        if f.get("references"):
            doc.add_paragraph("References:", style="Heading 3")
            for r in f["references"]:
                doc.add_paragraph(str(r), style="List Bullet")

    doc.add_heading("Remediation Roadmap", level=1)
    buckets = {"critical": "Immediate (0-7 days)", "high": "Short term (1-30 days)",
               "medium": "Medium term (30-90 days)", "low": "Long term (90+ days)",
               "info": "Long term (90+ days)"}
    grouped = {}
    for f in findings:
        grouped.setdefault(buckets.get(f.get("severity", "medium"), "Long term (90+ days)"), []).append(f)
    for bucket in ("Immediate (0-7 days)", "Short term (1-30 days)",
                   "Medium term (30-90 days)", "Long term (90+ days)"):
        doc.add_heading(bucket, level=2)
        items = grouped.get(bucket, [])
        if not items:
            doc.add_paragraph("(none)")
        for f in items:
            doc.add_paragraph("[" + str(f.get("id", "")) + "] " + str(f.get("title", "")), style="List Bullet")

    doc.save(str(out_path))
    return True


def _build_meta(client: str, engagement: str, author: str,
                scope: List[str], date_start: str, date_end: str) -> Dict:
    return {
        "title": "Security Assessment Report",
        "client": client or "(client name)",
        "engagement": engagement or "External Penetration Test",
        "date_start": date_start or datetime.utcnow().strftime("%Y-%m-%d"),
        "date_end": date_end or datetime.utcnow().strftime("%Y-%m-%d"),
        "author": author or "(author)",
        "report_id": "RS-" + datetime.utcnow().strftime("%Y%m%d-%H%M%S"),
        "scope": scope or [],
        "evidence_files": "(attached)",
    }


def cmd_markdown(client, engagement, author, scope, date_start, date_end, out_file):
    meta = _build_meta(client, engagement, author, scope, date_start, date_end)
    md = _render_md(meta)
    out = Path(out_file) if out_file else DELIV_DIR / ("report_" + meta["report_id"] + ".md")
    out.write_text(md)
    print_ok("wrote " + str(out))
    print_kv("bytes", str(len(md)))
    print_kv("findings", str(len(_load_findings())))
    return 0


def cmd_html(client, engagement, author, scope, date_start, date_end, out_file):
    meta = _build_meta(client, engagement, author, scope, date_start, date_end)
    doc = _render_html(meta)
    out = Path(out_file) if out_file else DELIV_DIR / ("report_" + meta["report_id"] + ".html")
    out.write_text(doc)
    print_ok("wrote " + str(out))
    print_kv("bytes", str(len(doc)))
    print_kv("findings", str(len(_load_findings())))
    return 0


def cmd_docx(client, engagement, author, scope, date_start, date_end, out_file):
    meta = _build_meta(client, engagement, author, scope, date_start, date_end)
    out = Path(out_file) if out_file else DELIV_DIR / ("report_" + meta["report_id"] + ".docx")
    ok = _render_docx(meta, out)
    if ok:
        print_ok("wrote " + str(out))
    return 0 if ok else 1


def run_cli(args):
    p = argparse.ArgumentParser(prog="redsky report deliverable", add_help=False)
    p.add_argument("-h", "--help", action="store_true")
    p.add_argument("action", nargs="?", default="markdown",
                   choices=["markdown", "html", "docx"])
    p.add_argument("--client", default="")
    p.add_argument("--engagement", default="")
    p.add_argument("--author", default="")
    p.add_argument("--scope", action="append", default=[])
    p.add_argument("--date-start", default="")
    p.add_argument("--date-end", default="")
    p.add_argument("--out", default="")

    try:
        ns = p.parse_args(args)
    except SystemExit:
        print_err("usage: redsky report deliverable <markdown|html|docx> [opts]")
        return 2

    if ns.help:
        print_info("markdown --client 'Acme' --engagement 'External PT' --author 'nono' \\")
        print_info("         --scope 10.0.0.0/24 --scope app.acme.com --out report.md")
        print_info("html     self-contained HTML")
        print_info("docx     Word (python-docx needed)")
        return 0

    if ns.action == "markdown":
        return cmd_markdown(ns.client, ns.engagement, ns.author, ns.scope,
                            ns.date_start, ns.date_end, ns.out)
    if ns.action == "html":
        return cmd_html(ns.client, ns.engagement, ns.author, ns.scope,
                        ns.date_start, ns.date_end, ns.out)
    if ns.action == "docx":
        return cmd_docx(ns.client, ns.engagement, ns.author, ns.scope,
                        ns.date_start, ns.date_end, ns.out)
    return 2


if __name__ == "__main__":
    sys.exit(run_cli(sys.argv[1:]))
